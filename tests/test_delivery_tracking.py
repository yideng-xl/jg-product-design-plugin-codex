"""Exercise header changes and customer mapping checks with synthetic data."""
import importlib.util
from pathlib import Path
import tempfile
import unittest
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('delivery_inspector', ROOT / 'skills/update-delivery-map/scripts/inspect_tracking.py')
inspector = importlib.util.module_from_spec(spec)
spec.loader.exec_module(inspector)


class TrackingTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / 'tracking.xlsx'

    def tracking(self, headers, rows):
        wb = openpyxl.Workbook()
        wb.active.append(headers)
        for row in rows:
            wb.active.append(row)
        wb.save(self.path)

    def test_inserted_classification_and_reordered_columns(self):
        headers = ['分类', '客户', '区域', *inspector.FIELDS]
        row = ['升级', '甲地调', '甲省', '正式', '合同类', None, None, None, None, None]
        self.tracking(headers, [row, ['交付', *row[1:]]])
        result = inspector.inspect(self.path, '网管')
        self.assertEqual(result['errors'], [])
        self.assertEqual((result['row_count'], result['unique_customers']), (2, 1))
        self.assertTrue(result['repeated_customers'][0]['same_map_fields'])
        self.assertEqual(result['classifications'], {'升级': 1, '交付': 1})

    def test_conflicting_customer_is_an_error(self):
        row = ['甲省', '甲地调', '正式', '合同类', '3.实施与验证', '进行中', None, None, None]
        changed = list(row)
        changed[5] = '取消'
        self.tracking(['区域', '客户', *inspector.FIELDS], [row, changed])
        result = inspector.inspect(self.path, '网管')
        self.assertIn({'conflicting_customer': '甲省 / 甲地调'}, result['errors'])

    def test_missing_column_and_blank_customer(self):
        self.tracking(['区域', '客户'], [['甲省', '甲地调']])
        self.assertTrue(inspector.inspect(self.path, '网管')['errors'])
        self.tracking(['区域', '客户', *inspector.FIELDS], [['甲省', None, '正式']])
        self.assertEqual(inspector.inspect(self.path, '网管')['errors'][0]['row'], 2)

    def test_jiguan_one_project_in_two_sheets(self):
        headers = ['地区', '客户', '标杆项目', *inspector.FIELDS]
        self.tracking(headers, [['甲省', '省调+超高压', '是', '正式', '合同类', '4.汇报与验收', '实施完成']])
        backbone = Path(self.tmp.name) / 'backbone.xlsx'
        wb = openpyxl.Workbook()
        wb.active.title = '国网三级调度'
        wb.create_sheet('超高压')
        for ws in wb:
            ws.append(['原客户（进度表 A+B 栏）'])
            ws.append(['甲省 / 省调+超高压'])
        wb['国网三级调度'].append(['乙省 / 省调'])
        wb.save(backbone)
        result = inspector.inspect(self.path, '集管', backbone)
        self.assertEqual(result['errors'], [])
        self.assertEqual(result['unmapped_customers'], [])
        self.assertEqual(result['historical_customers_absent'], ['乙省 / 省调'])
        self.assertEqual(len(result['multi_row_mappings']['甲省 / 省调+超高压']), 2)


if __name__ == '__main__':
    unittest.main()
