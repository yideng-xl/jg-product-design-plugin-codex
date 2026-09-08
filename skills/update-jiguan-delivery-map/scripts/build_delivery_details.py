#!/usr/bin/env python3
"""Join progress projects to deployment sites without inheriting environment fields."""
import argparse
import datetime as dt
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

import openpyxl

CORE = {'scheme': '部署方案', 'version': '集管版本', 'os': '操作系统',
        'redis': 'Redis 版本', 'database': '数据库及版本'}
PROGRESS_COLUMNS = ['地区', '客户', '交付阶段', '交付状态']
ENV_COLUMNS = {'method': '部署方式', 'scheme': '部署方案', 'version': '集管版本（最新）',
               'os': '操作系统', 'kernel': '系统内核版本', 'redis': 'Redis版本',
               'database': '数据库及版本', 'hardening': '加固及主机防火墙',
               'updatedAt': '最后修改时间'}


def clean(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat(sep=' ') if isinstance(value, dt.datetime) else value.isoformat()
    return str(value).strip() if value is not None and str(value).strip() else None


def header(value):
    return re.sub(r'^[①②③④⑤⑥⑦⑧⑨⑩\s]+', '', str(value or '')).strip()


def load_table(path, sheet, required):
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f'缺少工作表：{sheet}')
    ws = wb[sheet]
    columns = [header(c.value) for c in ws[1]]
    if len(columns) != len(set(columns)):
        raise ValueError(f'{sheet} 存在重复表头')
    missing = set(required) - set(columns)
    if missing:
        raise ValueError(f'{sheet} 缺少列：{sorted(missing)}')
    rows = [(i, dict(zip(columns, map(clean, row))))
            for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2)
            if any(v is not None for v in row)]
    wb.close()
    return rows


def token(value):
    return re.sub(r'[\s/\-－—_]', '', value or '').replace('分部网调', '网调').replace('分部', '网调')


def sites(region, customer):
    # The source's explicit '+' is the only rule that splits a project into sites.
    return [part if part.startswith(region) else region + part
            for part in customer.split('+')]


def source(path, sheet):
    return {'file': path.name, 'sheet': sheet, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}


def build(tracking, inventory, month):
    if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', month):
        raise ValueError('进度月份必须为 YYYY-MM')
    progress = load_table(tracking, '进度跟踪表', PROGRESS_COLUMNS)
    environments = load_table(inventory, '集管信息一览表', ['交付项目（客户）', *ENV_COLUMNS.values()])
    by_site, projects = {}, {}
    for row, rec in progress:
        region, customer = rec['地区'], rec['客户']
        if not region or not customer:
            raise ValueError(f'进度跟踪表第 {row} 行缺少地区或客户')
        key = f'{region} / {customer}'
        if key in projects:
            raise ValueError(f'重复项目键：{key}')
        projects[key] = rec
        for name in sites(region, customer):
            identifier = token(name)
            if identifier in by_site:
                raise ValueError(f'重复交付点：{name}')
            by_site[identifier] = {
                'site': name, 'region': region, 'projectKey': key,
                'stage': rec['交付阶段'], 'status': rec['交付状态'],
                'progressRow': row, 'inventoryRow': None,
                'combinedProgress': '+' in customer,
                **{k: None for k in ENV_COLUMNS},
            }
    seen, unmatched = set(), []
    for row, rec in environments:
        name = rec['交付项目（客户）']
        if not name:
            raise ValueError(f'集管信息一览表第 {row} 行缺少客户')
        key = token(name)
        if key in seen:
            raise ValueError(f'重复环境交付点：{name}')
        seen.add(key)
        if key not in by_site:
            unmatched.append({'site': name, 'inventoryRow': row})
            continue
        by_site[key].update({k: rec[col] for k, col in ENV_COLUMNS.items()})
        by_site[key].update({'inventoryRow': row, 'inventorySite': name})
    if unmatched:
        raise ValueError(f'环境表交付点未匹配，需核实别名，未生成输出：{unmatched}')
    records = list(by_site.values())
    for rec in records:
        rec['missing'] = [label for field, label in CORE.items() if not rec[field]]
        rec['notes'] = []
        if rec['combinedProgress']:
            rec['notes'].append('阶段及状态来自合并项目记录；环境按本交付点记录。')
        if not rec['inventoryRow']:
            rec['notes'].append('环境表未登记，待补收。')
        if rec['kernel'] and ('下架' in rec['kernel'] or '未部署' in rec['kernel']):
            rec['notes'].append('原表内核栏记录“' + rec['kernel'] + '”；当前运行环境待确认。')
        if rec['version'] and rec['hardening'] == '未部署':
            rec['notes'].append('版本已填写，加固栏记为未部署，当前部署情况待确认。')
    dates = [r['updatedAt'] for r in records if r['updatedAt']]
    return {
        'progressMonth': month,
        'sources': {'progress': source(tracking, '进度跟踪表'), 'inventory': source(inventory, '集管信息一览表')},
        'inventoryLatestUpdate': max(dates) if dates else None,
        'projectCount': len(projects), 'siteCount': len(records), 'inventoryCount': len(environments),
        'missingInventory': [r['site'] for r in records if not r['inventoryRow']],
        'completeCoreCount': sum(not r['missing'] for r in records),
        'missingCoreCounts': {label: sum(not r[k] for r in records) for k, label in CORE.items()},
        'projectStatusCounts': dict(Counter(r['交付状态'] or '待补收' for r in projects.values())),
        'records': records,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--tracking', type=Path, required=True)
    parser.add_argument('--inventory', type=Path, required=True)
    parser.add_argument('--date', required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    result = build(args.tracking, args.inventory, args.date)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({k: v for k, v in result.items() if k not in ['records', 'sources']}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
