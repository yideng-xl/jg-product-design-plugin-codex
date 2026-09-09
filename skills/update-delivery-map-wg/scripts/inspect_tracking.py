#!/usr/bin/env python3
"""Read-only checks for monthly delivery tracking workbooks. Requires openpyxl."""
import argparse
from collections import Counter, defaultdict
import json
from pathlib import Path
import sys

import openpyxl

FIELDS = ['合同类型', '项目分类', '交付阶段', '交付状态',
          '交付负责人', '交付开启日期', '交付完成日期']


def clean(value):
    return str(value).strip() if value is not None else ''


def inspect(path, product, backbone=None, sheet=None):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.active
    headers = [clean(v) for v in next(ws.values)]
    errors = []
    duplicates = [h for h, n in Counter(headers).items() if h and n > 1]
    if duplicates:
        errors.append({'duplicate_headers': duplicates})
    region = '区域' if '区域' in headers else '地区'
    fields = [region, '客户', *FIELDS]
    if product == '集管':
        fields.insert(4, '标杆项目')
    missing = [h for h in fields if h not in headers]
    result = {'source': str(Path(path).resolve()), 'sheet': ws.title,
              'product': product, 'headers': headers, 'errors': errors}
    if missing:
        errors.append({'missing_columns': missing})
        wb.close()
        return result
    indices = {h: headers.index(h) for h in fields}
    groups = defaultdict(list)
    categories, stages, statuses = Counter(), Counter(), Counter()
    count = 0
    for number, values in enumerate(ws.values, 1):
        if number == 1 or not any(v is not None for v in values):
            continue
        record = {h: values[i] for h, i in indices.items()}
        if not clean(record[region]) or not clean(record['客户']):
            errors.append({'row': number, 'empty_region_or_customer': True})
            continue
        key = f"{clean(record[region])} / {clean(record['客户'])}"
        category = clean(values[headers.index('分类')]) if '分类' in headers else ''
        groups[key].append({'row': number, 'classification': category, 'fields': record})
        categories[category or '(空)'] += 1
        stages[clean(record['交付阶段']) or '(空)'] += 1
        statuses[clean(record['交付状态']) or '(空)'] += 1
        count += 1
    wb.close()
    repeated = []
    for key, records in groups.items():
        if len(records) < 2:
            continue
        same = all(r['fields'] == records[0]['fields'] for r in records[1:])
        repeated.append({'key': key, 'rows': [r['row'] for r in records],
                         'same_map_fields': same,
                         'classifications': [r['classification'] for r in records]})
        if not same:
            errors.append({'conflicting_customer': key})
    result.update(row_count=count, unique_customers=len(groups),
                  classifications=dict(categories), stages=dict(stages),
                  statuses=dict(statuses), repeated_customers=repeated)
    if backbone:
        bw = openpyxl.load_workbook(backbone, data_only=True, read_only=True)
        locations = defaultdict(list)
        expected = ['国网三级调度'] + (['超高压'] if product == '集管' else [])
        for sn in expected:
            if sn not in bw.sheetnames:
                errors.append({'missing_backbone_sheet': sn})
                continue
            bs = bw[sn]
            bh = [clean(v) for v in next(bs.values)]
            key_indices = [i for i, h in enumerate(bh) if h.startswith('原客户')]
            if len(key_indices) != 1:
                errors.append({'invalid_backbone_key_column': sn})
                continue
            ki = key_indices[0]
            for number, values in enumerate(bs.values, 1):
                if number > 1 and clean(values[ki]):
                    locations[clean(values[ki])].append(f'{sn}!R{number}')
        bw.close()
        result['unmapped_customers'] = sorted(set(groups) - set(locations))
        result['historical_customers_absent'] = sorted(set(locations) - set(groups))
        result['multi_row_mappings'] = {k: v for k, v in locations.items() if len(v) > 1}
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('tracking', type=Path)
    parser.add_argument('--product', required=True, choices=['网管', '集管'])
    parser.add_argument('--backbone', type=Path)
    parser.add_argument('--sheet')
    args = parser.parse_args()
    try:
        result = inspect(args.tracking, args.product, args.backbone, args.sheet)
    except (OSError, ValueError, KeyError) as exc:
        print(json.dumps({'errors': [str(exc)]}, ensure_ascii=False))
        return 2
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
    return 2 if result['errors'] else 0


if __name__ == '__main__':
    sys.exit(main())
